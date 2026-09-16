"""Practical text/document ingestion and reviewable architecture candidates."""
from __future__ import annotations

import io
import json
from pathlib import Path


SUPPORTED_EXTENSIONS = {".txt", ".md", ".docx", ".pdf"}


def extract_text(data: bytes, filename: str) -> str:
    """Extract text only; binary documents are never sent to an LLM."""
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Unsupported document type: {suffix or 'unknown'}")
    if suffix in {".txt", ".md"}:
        return data.decode("utf-8", errors="replace")
    if suffix == ".docx":
        from docx import Document
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        return "\n".join(parts)
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(data))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def extract_tables(data: bytes, filename: str) -> list[list[list[str]]]:
    """Return DOCX tables as plain values for deterministic labelled-field parsing."""
    if Path(filename).suffix.lower() != ".docx":
        return []
    from docx import Document
    doc = Document(io.BytesIO(data))
    return [[[cell.text for cell in row.cells] for row in table.rows]
            for table in doc.tables]


def candidate_prompt(text: str, known_applications: list[str]) -> str:
    schema = """Return ONLY a JSON array. Each item must have:
{"candidate_type":"application|interface|information_object|relationship",
"name":"","source_application":"","target_application":"","description":"",
"confidence":0.0}"""
    return (
        "Extract architecture candidates from the document. Do not invent IDs; "
        "use names exactly as written or leave fields blank. Candidates are "
        "suggestions only and require human review.\n" + schema +
        "\nKnown applications:\n" + ", ".join(known_applications[:300]) +
        "\nDOCUMENT TEXT:\n" + text[:50000]
    )


def parse_candidates(response: str) -> list[dict]:
    raw = response.strip()
    if "```" in raw:
        raw = raw.split("```")[1].removeprefix("json").strip()
    start, end = raw.find("["), raw.rfind("]")
    if start < 0 or end < start:
        return []
    try:
        value = json.loads(raw[start:end + 1])
    except json.JSONDecodeError:
        return []
    return [x for x in value if isinstance(x, dict)]
